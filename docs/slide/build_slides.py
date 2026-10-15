"""Dựng deck "AI Kaizen Copilot — Motion Analyzer" (≤ 15 slide) trên template DENSO Factory Hacks 2026 bằng python-pptx.

Cách làm (giữ chuẩn template):
- Mỗi slide mới = NHÂN BẢN một slide mẫu (nền gradient + logo DENSO/FPT, header, footer, chữ CONFIDENTIAL…).
  Slide mẫu khớp nội dung (bìa, tóm tắt, bài toán, giải pháp, hiệu quả, triển khai, kết luận, kết) được giữ nguyên
  bố cục và chỉ thay chữ (giữ font/màu/cỡ của template). Slide kết quả mới chỉ lấy KHUNG (nền, header, footer)
  rồi vẽ card/ô số liệu/ảnh theo đúng bảng màu template.
- Ảnh lấy từ run outputs/runs/20261013-003627_pred_mediapipe_dry_u001-2trials (đã chép vào docs/slide/assets/).
- Kịch bản thuyết trình của từng slide ghi vào speaker notes (đồng thời có trong script.md).

    python docs/slide/build_slides.py      # → docs/slide/AI_Kaizen_Copilot_Motion_Analyzer.pptx
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "AI Kaizen Copilot — DENSO Factory Hacks 2026.pptx"
ASSETS = HERE / "assets"
OUT = HERE / "AI_Kaizen_Copilot_Motion_Analyzer.pptx"

# Bảng màu & font lấy từ template
BLUE, NAVY, TEXT, MUTED = "1265A8", "153866", "1F2D3D", "586777"
CARD, BORDER, TINT, MINT, GREEN, GREEN2 = "F7F9FB", "DCE6EF", "EAF3FC", "E9F8F0", "30D158", "00B96B"
FONT, FONT_B = "Aptos", "Aptos Bold"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- nhân bản slide mẫu
def clone(prs, src, keep: set[str] | None = None):
    """Thêm slide mới chép shape (và quan hệ ảnh) từ slide mẫu `src`; keep = chỉ giữ shape có tên trong tập."""
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    rid = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
            continue
        rid[rel.rId] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                        else new.part.relate_to(rel.target_part, rel.reltype))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        name = el.xpath("./*[1]/p:cNvPr/@name")
        if keep is not None and (not name or name[0] not in keep):
            continue
        e = copy.deepcopy(el)
        for node in e.iter():
            for attr in ("embed", "id", "link"):
                v = node.get(R_NS + attr)
                if v in rid:
                    node.set(R_NS + attr, rid[v])
        tree.append(e)
    return new


def _find(container, name: str):
    for sh in container.shapes:
        if sh.name == name:
            return sh
        if isinstance(sh, GroupShape):
            r = _find(sh, name)
            if r is not None:
                return r
    return None


def find(container, name: str):
    """Tìm shape theo tên (kể cả trong group); KeyError nếu không có."""
    r = _find(container, name)
    if r is None:
        raise KeyError(name)
    return r


def text_in(group):
    """Shape có chữ đầu tiên trong group (vd nhãn của pill/band)."""
    for sh in group.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh
    raise KeyError(group.name)


def set_text(shape, lines: str | list[str]):
    """Thay chữ nhưng GIỮ định dạng run đầu tiên của template (font, màu, cỡ, giãn dòng)."""
    if isinstance(shape, GroupShape):
        shape = text_in(shape)
    lines = [lines] if isinstance(lines, str) else lines
    txBody = shape.text_frame._txBody
    ps = txBody.findall(qn("a:p"))
    p0 = ps[0]
    for p in ps[1:]:
        txBody.remove(p)
    for el in list(p0):
        if el.tag in (qn("a:br"), qn("a:fld")) or (el.tag == qn("a:r") and el is not p0.find(qn("a:r"))):
            p0.remove(el)
    proto = copy.deepcopy(p0)
    prev = p0
    for i, line in enumerate(lines):
        p = p0 if i == 0 else copy.deepcopy(proto)
        if i:
            prev.addnext(p)
            prev = p
        p.find(qn("a:r")).find(qn("a:t")).text = line


def delete(slide, *names):
    for n in names:
        el = find(slide, n)._element
        el.getparent().remove(el)


# ---------------------------------------------------------------- vẽ phần tử theo style template
def _runs(par, text: str, size: float, color: str, bold: bool):
    """**đậm** trong chuỗi → run đậm."""
    for k, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = par.add_run()
        r.text = seg
        b = bold or k % 2 == 1
        r.font.size, r.font.bold, r.font.name = Pt(size), b, FONT_B if b else FONT
        r.font.color.rgb = rgb(color)


def tb(slide, x, y, w, h, paras, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
       spacing=1.12, bullet=False, after=4):
    """Text box không padding. paras: list chuỗi (hỗ trợ **đậm**) hoặc (chuỗi, {size,color,bold})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, item in enumerate([paras] if isinstance(paras, (str, tuple)) else paras):
        txt, o = (item, {}) if isinstance(item, str) else item
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = o.get("align", align)
        par.line_spacing = spacing
        par.space_after = Pt(o.get("after", after))
        if o.get("bullet", bullet):
            pPr = par._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.25)))
            pPr.set("indent", str(-Inches(0.25)))
            bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
            bc = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bf)
            pPr.append(bc)
        _runs(par, txt, o.get("size", size), o.get("color", color), o.get("bold", bold))
    return box


def card(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.18):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.2)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def pill(slide, x, y, w, h, label, fill=TINT, color=BLUE, size=14):
    s = card(slide, x, y, w, h, fill=fill, line=None, radius=h / 2)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    par = tf.paragraphs[0]
    par.alignment = PP_ALIGN.CENTER
    _runs(par, label, size, color, True)
    return s


def circle(slide, x, y, d, label, fill=BLUE, color="FFFFFF", size=16):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _runs(tf.paragraphs[0], label, size, color, True)
    return s


def picture(slide, name: str, x, y, w, h, border=True, align="center"):
    """Ảnh vừa khung (giữ tỉ lệ), căn giữa / trái trong khung."""
    from PIL import Image
    p = ASSETS / name
    iw, ih = Image.open(p).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = x + (w - pw) / 2 if align == "center" else x
    pic = slide.shapes.add_picture(str(p), Inches(px), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    if border:
        pic.line.color.rgb = rgb(BORDER)
        pic.line.width = Pt(1)
    return pic


def stat(slide, x, y, w, h, number, label, note="", num_color=BLUE, num_size=40):
    card(slide, x, y, w, h)
    tb(slide, x + 0.3, y + 0.22, w - 0.6, 0.85, number, size=num_size, color=num_color, bold=True)
    tb(slide, x + 0.3, y + 1.05, w - 0.6, 0.5, label, size=17, color=NAVY, bold=True)
    if note:
        tb(slide, x + 0.3, y + 1.5, w - 0.6, h - 1.6, note, size=14, color=MUTED)


def notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text.strip()


# ---------------------------------------------------------------- khung slide nội dung (từ slide mẫu 4)
FRAME4 = {"Freeform 2", "TextBox 3", "Group 4", "Group 6", "Group 8", "Group 10", "Group 13", "TextBox 16", "TextBox 17"}


def content_slide(prs, T, title, subtitle, foot_title, foot_text):
    s = clone(prs, T[4], FRAME4)
    t = find(s, "TextBox 16")
    set_text(t, title)
    t.left, t.width = Inches(4.9), Inches(14.4)
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    set_text(find(s, "TextBox 17"), subtitle)
    set_text(find(s, "Group 10"), foot_title)
    set_text(find(s, "Group 13"), foot_text)
    return s


# ================================================================ NỘI DUNG
def build():
    prs = Presentation(str(TEMPLATE))
    T = {i + 1: s for i, s in enumerate(prs.slides)}
    n_template = len(prs.slides)

    # 1 ── Bìa
    s = clone(prs, T[1])
    set_text(find(s, "TextBox 6"), "AI KAIZEN COPILOT")
    set_text(find(s, "TextBox 7"), "Motion Analyzer: đo Muda · Mura · Muri từ video thao tác")
    set_text(find(s, "TextBox 8"), "A2 • TÍCH LŨY KNOW-HOW CHUYÊN GIA CHO AI HỌC")
    tb(s, 9.61, 10.1, 9.9, 0.4, "Đội thi: Trịnh Gia Huy (nhóm trưởng) · Vũ Văn Mạnh · Nguyễn Trần Quang · Trịnh Hoài Nam",
       size=14, color="FFFFFF")
    notes(s, """
[~45 giây — mở đầu]
Kính chào ban giám khảo. Chúng em là đội AI Kaizen Copilot, gồm Trịnh Gia Huy (nhóm trưởng), Vũ Văn Mạnh, Nguyễn Trần Quang và Trịnh Hoài Nam.
Đề tài của đội là "Motion Analyzer": biến video thao tác lắp ráp thành các chỉ số Lean Muda – Mura – Muri và đề xuất Kaizen có kiểm chứng.
Đề tài thuộc danh mục A2: tích lũy know-how của người thợ giỏi thành dữ liệu để AI học và để đào tạo người mới.
Trong khoảng 15 phút, chúng em sẽ trình bày bài toán, giải pháp, kết quả đo thật trên dataset HA4M và kế hoạch pilot tại DENSO.
""")

    # 2 ── Tóm tắt đề xuất (bảng 6 câu hỏi của BTC)
    s = clone(prs, T[2])
    tt = find(s, "TextBox 23")
    set_text(tt, "TÊN CHỦ ĐỀ:  AI KAIZEN COPILOT — MOTION ANALYZER")
    tt.width = Inches(17.5)
    answers = [
        "Đo đồng thời Muda–Mura–Muri từ camera RGB thường; “số do code tính, LLM chỉ giải thích” với validator tự động "
        "chặn số/trích dẫn bịa; kiểm chứng kép bằng nhãn chuyên gia HA4M và cảm biến Azure Kinect.",
        "Bấm giờ thủ công 2–4 giờ/trạm, phụ thuộc người đo, chỉ lấy vài chu kỳ; know-how người thợ giỏi là tri thức ngầm, "
        "mất khi nghỉ việc; lãng phí, biến động, quá tải không được đo liên tục.",
        "MediaPipe Pose trên CPU → lọc One Euro → zone/rule tách bước SOP → chỉ số 3M → 7 rule phát hiện → LLM "
        "(Gemini/Claude) + validator → dashboard. PoC đã chạy end-to-end; pilot 4 tuần tại 1 trạm.",
        "HA4M U001: sai số cycle time 0.6 s/77.9 s, F1@IoU0.5 bước 0.75, MAE bước 0.62 s; cổ tay phải 95% frame lệch "
        "< 50 px so với Azure Kinect; phân tích 2 chu kỳ ≈ 9 s sau khi có keypoint.",
        "Chỉ cần cấu hình trạm (zone + SOP) và KB riêng; mỗi lần chạy là một run có manifest để so sánh trước/sau, "
        "giữa trạm, giữa người; lưu golden cycle cho đào tạo.",
        "Computer vision & pose, chuỗi thời gian, Lean/IE (3M, ECRS, Yamazumi), LLM có kiểm chứng, dashboard "
        "Streamlit và kiểm thử tự động (30 test).",
    ]
    table = find(s, "Table 22").table
    for i, a in enumerate(answers, start=1):
        set_text(_CellProxy(table.cell(i, 1)), a)
    set_text(find(s, "Group 15"), "Nếu được chọn, đội sẽ pilot 4 tuần tại một trạm DENSO: quay video → dashboard "
                                  "Muda/Mura/Muri có bằng chứng → kỹ sư duyệt → đo lại trước/sau.")
    set_text(find(s, "Group 19"), "DENSO Factory Hacks 2026  |  AI Kaizen Copilot — Motion Analyzer")
    notes(s, """
[~1 phút — tóm tắt cho giám khảo]
Slide này trả lời nhanh 6 câu hỏi của ban tổ chức.
Điểm mới: chúng em đo cả 3M chứ không chỉ bấm giờ, và tách bạch "con số do thuật toán tính – LLM chỉ giải thích", có validator tự động chặn số liệu hay trích dẫn bịa.
Vấn đề: bấm giờ thủ công tốn 2 đến 4 giờ một trạm và know-how của người thợ giỏi không được lưu lại.
Kết quả đã đo: sai số cycle time 0,6 giây trên chu kỳ 77,9 giây, F1 phân đoạn bước 0,75, và cổ tay phải khớp cảm biến Azure Kinect ở 95% khung hình.
Thông điệp chính: nếu được chọn, đội sẽ pilot 4 tuần tại một trạm, đo trước và sau bằng cùng một phương pháp.
""")

    # 3 ── Team (khung slide mẫu 3, 4 thành viên)
    s = clone(prs, T[3], {"Freeform 2", "TextBox 3", "Group 4", "Group 6", "TextBox 8", "Group 9", "Group 11",
                          "Group 14", "Group 50", "Group 52"})
    members = [("TH", "Trịnh Gia Huy", "Nhóm trưởng"), ("VM", "Vũ Văn Mạnh", "Thành viên"),
               ("NQ", "Nguyễn Trần Quang", "Thành viên"), ("TN", "Trịnh Hoài Nam", "Thành viên")]
    x0, w, gap = 0.85, 4.35, 0.32
    for k, (ini, name, role) in enumerate(members):
        x = x0 + k * (w + gap)
        card(s, x, 2.45, w, 3.25)
        circle(s, x + w / 2 - 0.65, 2.75, 1.3, ini, fill=BLUE if k else NAVY, size=26)
        tb(s, x + 0.2, 4.25, w - 0.4, 0.5, name, size=23, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        pill(s, x + w / 2 - 1.0, 4.95, 2.0, 0.45, role, fill=MINT if k == 0 else TINT,
             color=GREEN2 if k == 0 else BLUE, size=14)
    tb(s, 0.85, 6.05, 18.3, 0.5, "Các mảng đội đã xây dựng trong PoC", size=20, color=NAVY, bold=True)
    areas = [("Computer Vision & Motion", "Trích keypoint MediaPipe, lọc nhiễu, zone, sự kiện thao tác"),
             ("Kaizen Metrics & Rules", "Phân đoạn bước SOP, chỉ số Muda/Mura/Muri, 7 rule phát hiện"),
             ("LLM Copilot & Validator", "Giải thích có trích dẫn KB/SOP, chặn số bịa, retry có phản hồi"),
             ("Dashboard & Đánh giá", "Streamlit 8 tab, video overlay, so sánh run, 30 test tự động")]
    for k, (h, d) in enumerate(areas):
        x = x0 + k * (w + gap)
        card(s, x, 6.65, w, 1.7, fill="FFFFFF")
        tb(s, x + 0.25, 6.85, w - 0.5, 0.45, h, size=17, color=BLUE, bold=True)
        tb(s, x + 0.25, 7.35, w - 0.5, 1.1, d, size=14.5, color=TEXT)
    set_text(find(s, "Group 52"), "PHỐI HỢP: CV & MOTION  •  KAIZEN METRICS  •  LLM + VALIDATOR  •  DASHBOARD & ĐÁNH GIÁ")
    set_text(find(s, "Group 14"), "Nhóm trưởng Trịnh Gia Huy điều phối kỹ thuật và trình bày; các mảng do cả đội cùng xây dựng.")
    notes(s, """
[~45 giây — giới thiệu đội]
Đội gồm 4 thành viên: Trịnh Gia Huy là nhóm trưởng, cùng Vũ Văn Mạnh, Nguyễn Trần Quang và Trịnh Hoài Nam.
Trong giai đoạn hackathon, đội đã xây dựng đủ 4 mảng của một hệ thống chạy thật:
thứ nhất là thị giác máy tính để trích khung xương và sự kiện thao tác;
thứ hai là lớp chỉ số Kaizen – phân đoạn bước SOP và tính Muda, Mura, Muri;
thứ ba là Kaizen Copilot dùng LLM có kiểm chứng;
và cuối cùng là dashboard cùng bộ kiểm thử tự động 30 test.
""")

    # 4 ── Bài toán (slide mẫu 4)
    s = clone(prs, T[4])
    set_text(find(s, "TextBox 17"), "Know-how thao tác của người thợ giỏi chưa được đo, chưa được lưu và khó chuyển giao.")
    for t_name, b_name, title, body in [
        ("TextBox 27", "TextBox 28", "TRI THỨC NGẦM DỄ MẤT",
         "Thứ tự động tác, nhịp làm việc, phối hợp hai tay chỉ nằm trong cơ thể người thợ giỏi; truyền miệng qua kèm "
         "cặp và mất khi nhân sự thay đổi."),
        ("TextBox 38", "TextBox 39", "BẤM GIỜ THỦ CÔNG",
         "Kỹ sư IE đo 2–4 giờ cho một trạm, ghi tay, mỗi người đo một kiểu; chỉ lấy được vài chu kỳ nên bỏ sót biến động."),
        ("TextBox 49", "TextBox 50", "3M CHƯA ĐƯỢC ĐO",
         "Lãng phí (Muda), biến động (Mura), quá tải cổ tay (Muri) không được theo dõi hằng ngày; đề xuất Kaizen khó "
         "chứng minh bằng số liệu trước–sau."),
    ]:
        set_text(find(s, t_name), title)
        set_text(find(s, b_name), body)
    set_text(find(s, "Group 53"), "HƯỚNG GIẢI:  camera đo mọi chu kỳ + LLM giải thích có kiểm chứng + kỹ sư IE duyệt.")
    set_text(find(s, "Group 10"), "BỐI CẢNH")
    set_text(find(s, "Group 13"), "Thời gian bấm giờ 2–4 giờ/trạm là ước tính từ thực tế IE, cần xác nhận tại DENSO.")
    notes(s, """
[~1 phút — bài toán]
Ở dây chuyền lắp ráp, cái làm nên năng suất của người thợ giỏi là thứ tự động tác, nhịp làm việc và cách phối hợp hai tay. Đây là tri thức ngầm: không nằm trong tài liệu, chỉ truyền qua kèm cặp, và mất khi người thợ nghỉ.
Muốn đo thì hiện nay kỹ sư IE phải bấm giờ thủ công, 2 đến 4 giờ cho một trạm, mỗi người đo một kiểu và chỉ lấy được vài chu kỳ.
Hệ quả là ba loại lãng phí theo ngôn ngữ Lean – Muda, Mura, Muri – không được theo dõi liên tục, và đề xuất Kaizen khó chứng minh bằng số liệu trước sau.
Hướng giải của đội: dùng camera đo mọi chu kỳ, dùng LLM giải thích theo tri thức Kaizen nhưng có kiểm chứng, và kỹ sư là người quyết định.
""")

    # 5 ── Luồng giải pháp (slide mẫu 5)
    s = clone(prs, T[5])
    t = find(s, "TextBox 8")
    set_text(t, "LUỒNG GIẢI PHÁP")
    set_text(find(s, "TextBox 17"), "Camera → khung xương → bước SOP → chỉ số 3M → findings → giải thích có trích dẫn "
                                    "→ kỹ sư duyệt.")
    for grp, t_name, b_name, lab, title, body in [
        ("Group 22", "TextBox 25", "TextBox 26", "OBSERVE", "Video → khung xương",
         "MediaPipe Pose chạy trên CPU, lọc nhiễu One Euro; chỉ lưu keypoint, không cần nhận diện khuôn mặt."),
        ("Group 34", "TextBox 37", "TextBox 38", "MEASURE", "Bước SOP & 3M",
         "Zone + rule tách 12 bước; tính cycle time, %VA, thời gian chờ, số lần lấy, CV, góc cổ tay."),
        ("Group 46", "TextBox 49", "TextBox 50", "EXPLAIN", "Findings + LLM",
         "7 rule tìm vấn đề, xếp hạng theo ảnh hưởng; LLM giải thích có trích dẫn KB/SOP, validator chặn số bịa."),
        ("Group 58", "TextBox 61", "TextBox 62", "VERIFY", "Kỹ sư duyệt",
         "Dashboard + video bằng chứng; duyệt hoặc bác bỏ, thử cải tiến, đo lại cùng phương pháp trước–sau."),
    ]:
        set_text(find(s, grp), lab)
        set_text(find(s, t_name), title)
        set_text(find(s, b_name), body)
    set_text(find(s, "Group 65"), "ĐẦU RA: Gantt bước • Yamazumi & %VA • Mura • Muri • findings có video bằng chứng • dashboard")
    set_text(find(s, "Group 14"), "Mọi con số do thuật toán tính và truy vết được; LLM chỉ giải thích; mọi đề xuất cần "
                                  "kỹ sư duyệt.")
    notes(s, """
[~1 phút — giải pháp]
Giải pháp là một luồng khép kín bốn khâu.
OBSERVE: camera thường quay trạm, MediaPipe trích khung xương ngay trên CPU, chỉ lưu tọa độ khớp nên không cần giữ khuôn mặt.
MEASURE: hệ thống tách 12 bước SOP dựa trên vùng khay linh kiện và quy tắc, rồi tính cycle time, tỷ lệ giá trị gia tăng, thời gian chờ, số lần lấy linh kiện, độ biến động và góc cổ tay.
EXPLAIN: 7 quy tắc tìm vấn đề, xếp hạng theo ảnh hưởng; LLM giải thích có trích dẫn knowledge base và SOP.
VERIFY: kỹ sư xem dashboard và video bằng chứng, duyệt hay bác bỏ, rồi đo lại bằng đúng phương pháp đó.
""")

    # 6 ── Kiến trúc pipeline
    s = content_slide(prs, T, "KIẾN TRÚC PIPELINE",
                      "9 bước độc lập, mỗi bước đọc file → ghi file; mỗi lần chạy là một run có manifest để tái lập "
                      "và so sánh.",
                      "KIẾN TRÚC", "Mỗi bước chạy lại độc lập được; 30 test tự động.")
    stages = [("00", "Inspect", "chốt schema dữ liệu thật", BLUE), ("01", "Keypoints", "MediaPipe trên CPU", BLUE),
              ("02", "Features", "lọc nhiễu, tốc độ, zone", BLUE), ("03", "Events", "lấy · lắp · chờ · di chuyển", BLUE),
              ("04", "Steps", "oracle | pred + đánh giá", BLUE), ("05", "Metrics", "Muda · Mura · Muri", GREEN2),
              ("06", "Findings", "rule R1–R7, xếp hạng", GREEN2), ("07", "Explain", "LLM + validator", NAVY),
              ("08", "Report", "Markdown, PNG, video", NAVY)]
    bw, bg, bx = 1.84, 0.2, 0.9
    for k, (num, name, desc, col) in enumerate(stages):
        x = bx + k * (bw + bg)
        card(s, x, 2.85, bw, 2.45, fill=CARD)
        circle(s, x + bw / 2 - 0.35, 3.05, 0.7, num, fill=col, size=17)
        tb(s, x + 0.12, 3.9, bw - 0.24, 0.45, name, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        tb(s, x + 0.12, 4.38, bw - 0.24, 0.85, desc, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    for lab, x0_, x1_, col, fill in [("THỊ GIÁC MÁY TÍNH", 0, 4, BLUE, TINT), ("CHỈ SỐ KAIZEN", 5, 6, GREEN2, MINT),
                                     ("COPILOT & BÁO CÁO", 7, 8, NAVY, TINT)]:
        xa = bx + x0_ * (bw + bg)
        pill(s, xa, 5.45, (x1_ - x0_ + 1) * bw + (x1_ - x0_) * bg, 0.42, lab, fill=fill, color=col, size=13)
    for k, (h, body) in enumerate([
        ("Số do code tính — LLM chỉ giải thích",
         "Mỗi chỉ số trong metrics.json kèm đơn vị và định nghĩa. Validator: trích dẫn phải có trong KB/SOP, mọi số "
         "phải có trong finding; sai thì hỏi lại LLM 1 lần kèm lỗi cụ thể."),
        ("Chế độ --oracle",
         "Dùng nhãn chuyên gia thay cho bước CV để chạy nửa sau độc lập. Chênh lệch oracle vs pred chính là sai số "
         "của phần thị giác máy tính."),
        ("Run có manifest",
         "outputs/runs/<run_id>/ lưu lệnh gốc, git hash, bản chụp cấu hình trạm. So sánh oracle vs pred, Gemini vs "
         "Claude, trước vs sau Kaizen bằng một lệnh."),
    ]):
        x = 0.9 + k * 6.1
        card(s, x, 6.1, 5.85, 2.55, fill="FFFFFF")
        tb(s, x + 0.3, 6.3, 5.25, 0.5, h, size=18, color=BLUE, bold=True)
        tb(s, x + 0.3, 6.85, 5.25, 1.75, body, size=15, color=TEXT)
    tb(s, 0.9, 8.9, 18.2, 0.45, "**Công nghệ:** Python · MediaPipe · NumPy/Pandas · Shapely · One Euro filter · "
                                "Gemini 3.8 Flash / Claude · Streamlit · Plotly · H.264 (imageio-ffmpeg)",
       size=15, color=MUTED)
    notes(s, """
[~1 phút — kiến trúc]
Về kỹ thuật, pipeline gồm 9 bước, mỗi bước là một script độc lập đọc file và ghi file, nên chạy lại hay thay một khâu rất dễ.
Năm bước đầu là thị giác máy tính: chốt cấu trúc dữ liệu thật, trích keypoint, lọc nhiễu, sinh sự kiện thao tác và tách bước.
Bước 5 và 6 là lớp Kaizen: tính chỉ số 3M và chạy quy tắc phát hiện. Bước 7 là Copilot với LLM, bước 8 sinh báo cáo.
Có ba nguyên tắc quan trọng: thứ nhất, mọi con số do code tính, LLM chỉ giải thích và bị validator kiểm tra.
Thứ hai, chế độ oracle dùng nhãn chuyên gia để kiểm tra nửa sau độc lập với sai số CV.
Thứ ba, mỗi lần chạy là một run có manifest, giúp tái lập và so sánh trước – sau.
""")

    # 7 ── Dữ liệu & độ tin cậy keypoint
    s = content_slide(prs, T, "DỮ LIỆU & ĐỘ TIN CẬY KEYPOINT",
                      "HA4M: 41 người lắp bộ bánh răng hành tinh, 12 bước, nhãn chuyên gia theo từng khung hình.",
                      "DỮ LIỆU", "Kinect cũng là ước lượng — dùng để đối chứng, không phải ground truth tuyệt đối.")
    picture(s, "frame_V001_step3.png", 0.9, 2.7, 8.9, 6.1)
    tb(s, 0.9, 8.88, 8.9, 0.5, "Khung hình overlay từ video demo của run (U001_V001, t = 17.6 s, bước 3.1): zone khay "
                               "tô theo màu bước, vệt cổ tay ~1 s, trạng thái hai tay.", size=13, color=MUTED)
    stat(s, 10.2, 2.7, 4.35, 2.95, "2 chu kỳ", "U001_V001 · U001_V002",
         "5.082 khung hình, 30 fps, 2048×1536; trích keypoint ≈ 7.5 fps trên CPU")
    stat(s, 14.8, 2.7, 4.35, 2.95, "95–96%", "Khung hình phát hiện người", "detect_rate 0.953 / 0.961")
    stat(s, 10.2, 5.85, 4.35, 2.95, "24.5 px", "Lệch trung vị cổ tay phải", "So với skeleton Azure Kinect (U001_V001)")
    stat(s, 14.8, 5.85, 4.35, 2.95, "95%", "Khung hình lệch < 50 px", "Tay phải. Tay trái hay bị thân che: P90 151.6 px "
                                                                      "→ gắn cờ kém tin cậy", num_color=GREEN2)
    notes(s, """
[~1 phút — dữ liệu]
Để phát triển và đánh giá khách quan, đội dùng HA4M, một dataset công khai: 41 người lắp bộ bánh răng hành tinh gồm 12 bước, có nhãn chuyên gia theo từng khung hình. Trong slide này là hai chu kỳ lắp của người U001.
Hình bên trái là một khung hình từ video demo do hệ thống sinh ra: vùng khay được tô theo màu bước, vệt cổ tay trong khoảng 1 giây và trạng thái hai tay.
Độ tin cậy keypoint: hơn 95% khung hình phát hiện được người. So với skeleton của cảm biến Azure Kinect, cổ tay phải lệch trung vị 24,5 pixel và 95% khung hình dưới 50 pixel.
Tay trái hay bị thân che nên kém tin cậy hơn; hệ thống gắn cờ điều này thay vì che giấu.
""")

    # 8 ── Phân đoạn bước vs nhãn
    s = content_slide(prs, T, "PHÂN ĐOẠN BƯỚC: CV vs NHÃN",
                      "Bước bắt đầu ở lần lấy đầu tiên tại khay của bước; so trực tiếp với nhãn chuyên gia.",
                      "ĐÁNH GIÁ", "Baseline rule-based, chưa huấn luyện model; v2: DTW / change-point.")
    picture(s, "gantt_U001_V001.png", 0.9, 2.62, 11.2, 3.4)
    picture(s, "gantt_U001_V002.png", 0.9, 6.12, 11.2, 3.4)
    card(s, 12.5, 2.62, 6.65, 4.55)
    tb(s, 12.8, 2.82, 6.1, 0.45, "Sai số so với nhãn", size=19, color=NAVY, bold=True)
    cols = [(15.2, "V001"), (17.2, "V002")]
    for x, lab in cols:
        tb(s, x, 3.42, 1.8, 0.4, lab, size=16, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    for r, (lab, v1, v2) in enumerate([("Sai số cycle time", "0.6 s", "2.27 s"), ("F1@IoU0.5 bước", "0.75", "0.714"),
                                       ("MAE thời lượng bước", "0.62 s", "0.92 s")]):
        y = 3.95 + r * 1.0
        tb(s, 12.8, y + 0.12, 2.4, 0.5, lab, size=15, color=TEXT)
        for (x, _), v in zip(cols, (v1, v2)):
            tb(s, x, y, 1.8, 0.6, v, size=27, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    card(s, 12.5, 7.37, 6.65, 2.15, fill=TINT, line=None)
    tb(s, 12.8, 7.55, 6.1, 1.9, ["**Giới hạn hiện tại:** bước 9 không có khay riêng → gộp vào “chưa gán” (xám).",
                                 "V002 chưa bắt được bước 3.2 và 5 — hệ thống báo thiếu bước thay vì đoán.",
                                 "Precision 0.80 / 0.91 · recall 0.71 / 0.59 (V001 / V002)."],
       size=14.5, color=TEXT, after=6)
    notes(s, """
[~1 phút 15 — kết quả phân đoạn]
Đây là kết quả quan trọng nhất về độ chính xác. Mỗi biểu đồ có hàng trên là bước theo nhãn chuyên gia, hàng dưới là bước do hệ thống tự tách, cùng hoạt động của hai tay.
Với chu kỳ V001, sai số cycle time chỉ 0,6 giây trên 77,9 giây, tức khoảng 0,8%; F1 phân đoạn bước 0,75; sai số thời lượng mỗi bước trung bình 0,62 giây.
Chu kỳ V002 khó hơn: F1 0,71, sai số cycle time 2,27 giây, vì hệ thống chưa bắt được hai lần lấy linh kiện.
Đây mới là baseline dựa trên quy tắc, chưa huấn luyện model nào. Điểm quan trọng là hệ thống trung thực về giới hạn: phần không chắc chắn được tô xám và báo thiếu bước, không tự đoán.
""")

    # 9 ── Muda · Mura · Muri
    s = content_slide(prs, T, "MUDA · MURA · MURI",
                      "Yamazumi theo bước, biến động giữa hai lần lắp và sàng lọc tư thế cổ tay — số do pipeline tính.",
                      "3M", "Rule bước chậm (median + 2·MAD) tự kích hoạt khi có ≥ 3 trial.")
    picture(s, "yamazumi.png", 0.9, 2.62, 6.4, 6.9)
    for k, (tag, col, fill, title, lines) in enumerate([
        ("MUDA", BLUE, TINT, "Lãng phí trong chu kỳ",
         ["%VA theo nhãn: **82.1%** (V001) · **85.3%** (V002)",
          "**15 / 14** lần lấy linh kiện mỗi chu kỳ; quãng đường cổ tay phải ≈ 13.5k px",
          "Phần xám trên Yamazumi = thời gian CV chưa gán được bước"]),
        ("MURA", GREEN2, MINT, "Biến động giữa hai lần lắp",
         ["Cycle time theo nhãn **77.9 s** vs **84.3 s** → CV **0.056**",
          "Bước biến động nhất (CV dự đoán): **B10 Cover CV 0.28**, **B12 siết vít CV 0.26**"]),
        ("MURI", NAVY, TINT, "Sàng lọc tư thế cổ tay (2D)",
         ["Cổ tay phải lệch > 30° trong **11.6% / 15.5%** thời gian (độ phủ 88–91%)",
          "Tay trái 23.1% nhưng độ phủ chỉ 27% → chỉ mang tính sàng lọc"]),
    ]):
        y = 2.62 + k * 2.35
        card(s, 7.7, y, 11.45, 2.15)
        pill(s, 7.95, y + 0.25, 1.35, 0.42, tag, fill=fill, color=col, size=14)
        tb(s, 9.5, y + 0.22, 9.4, 0.45, title, size=18, color=NAVY, bold=True)
        tb(s, 7.95, y + 0.85, 10.95, 1.25, lines, size=15, color=TEXT, bullet=True, after=3)
    notes(s, """
[~1 phút 15 — 3M]
Từ các bước đã tách, hệ thống tính đủ ba chiều Lean.
Muda – lãng phí: biểu đồ Yamazumi bên trái chồng thời lượng 12 bước cho mỗi chu kỳ; tỷ lệ giá trị gia tăng theo nhãn là 82 và 85%; mỗi chu kỳ có 15 và 14 lần lấy linh kiện.
Mura – biến động: cùng một người nhưng hai lần lắp chênh nhau gần 6,5 giây; bước đặt nắp và bước siết vít biến động nhiều nhất – đó là nơi nên chuẩn hóa thao tác trước.
Muri – quá tải: hệ thống sàng lọc góc cổ tay 2D; tay phải lệch quá 30 độ khoảng 12 đến 15% thời gian. Với tay trái, độ phủ chỉ 27% nên hệ thống ghi rõ là chỉ mang tính sàng lọc.
Khi nạp thêm dữ liệu từ 3 chu kỳ trở lên, quy tắc phát hiện bước chậm bất thường sẽ tự kích hoạt.
""")

    # 10 ── Findings & Copilot
    s = content_slide(prs, T, "FINDINGS & KAIZEN COPILOT",
                      "Rule tìm vấn đề và xếp hạng theo ảnh hưởng; LLM giải thích nhưng không được tạo số.",
                      "HUMAN-IN-THE-LOOP", "Mỗi finding kèm mốc thời gian + ảnh bằng chứng; kỹ sư IE duyệt hoặc bác bỏ.")
    picture(s, "evidence_F01.png", 0.9, 2.62, 8.3, 6.25)
    tb(s, 0.9, 8.95, 8.3, 0.5, "F01 · U001_V002 · t = 62.6 s · bước 11: tay phải ở khay vít/lục giác (bằng chứng do "
                               "08_report sinh).", size=13, color=MUTED)
    card(s, 9.6, 2.62, 9.55, 2.9)
    tb(s, 9.9, 2.8, 9.0, 0.45, "3 findings của run (xếp hạng)", size=18, color=NAVY, bold=True)
    for r, (fid, rule, txt, col) in enumerate([
        ("F01", "R2", "Lấy dư ở khay vít/lục giác: **4 lần**, SOP cần 3 (≈ 0.84 s/chu kỳ)", "E34948"),
        ("F02", "R4", "V002 thiếu bước 3.2 và 5 — giới hạn CV, độ tin cậy thấp", "EDA100"),
        ("F03", "R7", "Cổ tay trái lệch > 30° trong **23.1%** thời gian — sàng lọc 2D", "EDA100"),
    ]):
        y = 3.4 + r * 0.68
        pill(s, 9.9, y, 0.85, 0.42, fid, fill=col, color="FFFFFF", size=13)
        pill(s, 10.85, y, 0.7, 0.42, rule, fill=TINT, color=BLUE, size=13)
        tb(s, 11.75, y + 0.05, 7.2, 0.45, txt, size=15, color=TEXT)
    card(s, 9.6, 5.72, 9.55, 3.75, fill=TINT, line=None)
    tb(s, 9.9, 5.9, 9.0, 0.45, "Giải thích thật — Gemini 3.8 Flash (finding cổ tay)", size=17, color=NAVY, bold=True)
    tb(s, 9.9, 6.4, 9.0, 1.7, ["“Thời gian góc cổ tay left lệch > 30° chiếm 23.1%, vượt ngưỡng 20.0%… góc lệch P90 "
                               "là 67.4°.”",
                               "“Theo ECRS (Rearrange/Simplify), xem xét sắp xếp lại vị trí hoặc nghiêng khay linh "
                               "kiện… kỹ sư IE cần quan sát video để xác nhận.”"],
       size=14, color=TEXT, after=4)
    pill(s, 9.9, 8.25, 1.1, 0.4, "KB-13", fill="FFFFFF", color=BLUE, size=13)
    pill(s, 11.1, 8.25, 1.1, 0.4, "KB-01", fill="FFFFFF", color=BLUE, size=13)
    pill(s, 12.4, 8.25, 3.5, 0.4, "✓ Validator: đạt ngay lần 1", fill=MINT, color=GREEN2, size=13)
    tb(s, 9.9, 8.8, 9.0, 0.6, "Validator đã bắt lỗi thật (“phân vị 90” → vá whitelist) và chặn số bịa như “tiết kiệm "
                              "8.2 s” bằng retry có phản hồi.", size=13, color=MUTED)
    notes(s, """
[~1 phút 15 — findings & Copilot]
Từ các chỉ số, 7 quy tắc tìm ra vấn đề và xếp hạng theo ảnh hưởng nhân với độ tin cậy.
Finding đứng đầu là F01: ở chu kỳ V002, người thợ lấy linh kiện ở khay vít và lục giác 4 lần trong khi SOP chỉ cần 3. Ảnh bên trái là bằng chứng tự động: đúng thời điểm 62,6 giây, bước 11, tay phải đang ở khay vít.
Phần Copilot: LLM nhận finding và knowledge base Kaizen, viết quan sát, giả thuyết và đề xuất có trích dẫn. Đây là câu trả lời thật của Gemini 3.8 Flash, trích dẫn KB-13 về tư thế cổ tay và KB-01 về 3M.
Điểm mấu chốt là validator: mọi con số LLM viết phải có trong dữ liệu, mọi trích dẫn phải tồn tại. Trong thử nghiệm, validator đã chặn số bịa và yêu cầu LLM sửa – đúng nỗi lo "AI bịa" trong môi trường nhà máy.
""")

    # 11 ── Dashboard demo
    s = content_slide(prs, T, "DASHBOARD DEMO",
                      "Streamlit 8 tab, chỉ đọc kết quả của run — mở thẳng một run bằng đường link ?run=<run_id>.",
                      "DEMO", "Video overlay H.264 phát ngay trên trình duyệt.")
    picture(s, "ui_video_compare.png", 0.9, 2.62, 7.7, 6.9)
    picture(s, "ui_timeline.png", 8.95, 2.62, 10.2, 4.3)
    card(s, 8.95, 7.12, 10.2, 2.4)
    tb(s, 9.25, 7.3, 9.6, 0.45, "8 tab cho kỹ sư IE", size=18, color=NAVY, bold=True)
    tabs = ["Tổng quan", "Video & Timeline", "So sánh video", "Muda", "Mura & Muri", "Findings & Copilot",
            "So sánh run", "Chạy pipeline"]
    for k, lab in enumerate(tabs):
        pill(s, 9.25 + (k % 4) * 2.42, 7.88 + (k // 4) * 0.62, 2.28, 0.46, lab, size=13)
    notes(s, """
[~1 phút — demo dashboard; có thể chuyển sang demo trực tiếp]
Toàn bộ kết quả được đưa lên dashboard Streamlit gồm 8 tab, chỉ đọc kết quả của run nên không bao giờ tự tính lại hay sai lệch con số.
Bên trái là tab So sánh video: hai lần lắp V001 và V002 phát song song, bấm "Nhảy tới bước" thì cả hai video cùng về đầu một bước SOP để kỹ sư so cách thao tác.
Bên phải là timeline: hàng bước theo nhãn và hàng bước hệ thống dự đoán xếp chồng nhau để thấy ngay chỗ lệch.
Ngoài ra còn các tab Muda, Mura và Muri, Findings có trích dẫn bấm xem nguyên văn, so sánh run và nút chạy pipeline ngay từ giao diện.
""")

    # 12 ── Hiệu quả (slide mẫu 6)
    s = clone(prs, T[6])
    set_text(find(s, "TextBox 8"), "Hiệu quả mang lại")
    for grp, t_name, b_name, num, title, body in [
        ("Group 19", "TextBox 22", "TextBox 23", "0.6 s", "Sai số cycle time",
         "V001: 78.5 s dự đoán vs 77.9 s theo nhãn (≈ 0.8%)."),
        ("Group 26", "TextBox 29", "TextBox 30", "0.75", "F1 phân đoạn bước",
         "F1@IoU0.5 trên V001; MAE bước 0.62 s. V002: 0.714."),
        ("Group 33", "TextBox 36", "TextBox 37", "95%", "Cổ tay phải < 50 px",
         "So với skeleton Azure Kinect; trung vị lệch 24.5 px."),
        ("Group 40", "TextBox 43", "TextBox 44", "~9 s", "Phân tích 2 chu kỳ",
         "Bước 02→08 sau khi có keypoint; trích keypoint ~6 phút/video trên CPU."),
    ]:
        set_text(find(s, grp), num)
        set_text(find(s, t_name), title)
        set_text(find(s, b_name), body)
    set_text(find(s, "TextBox 47"), "BUSINESS CASE")
    set_text(find(s, "TextBox 48"), ["• Phần cứng gần như 0: webcam/điện thoại + laptop sẵn có, không cần GPU.",
                                     "• Đo 100% chu kỳ thay vì vài mẫu; thời gian người cho 1 lần time study ước tính "
                                     "< 30 phút (cần kiểm chứng).",
                                     "• Golden cycle của người thợ giỏi lưu thành chuẩn số hoá cho đào tạo."])
    set_text(find(s, "TextBox 52"), ["1. Baseline cùng trạm, cùng định nghĩa bước SOP.",
                                     "2. Đối chiếu nhãn / đánh giá của kỹ sư IE (như nhãn HA4M).",
                                     "3. Đo lại sau Kaizen bằng cùng pipeline; so sánh bằng compare_runs."])
    set_text(find(s, "Group 11"), "KẾT QUẢ ĐO TRÊN HA4M")
    set_text(find(s, "Group 14"), "Bốn số liệu trên đo thật trên HA4M U001 (V001, V002), baseline rule-based; hiệu quả "
                                  "vận hành tại DENSO là ước tính, cần pilot xác nhận.")
    notes(s, """
[~1 phút — hiệu quả]
Tóm lại các con số đã đo được – tất cả đo thật, không phải mục tiêu: sai số cycle time 0,6 giây, F1 phân đoạn bước 0,75, 95% khung hình cổ tay khớp cảm biến công nghiệp dưới 50 pixel, và phân tích hai chu kỳ chỉ mất khoảng 9 giây sau khi có keypoint.
Về kinh doanh: chi phí phần cứng gần như bằng 0 vì chỉ cần webcam và laptop; đo được 100% chu kỳ thay vì vài mẫu; và golden cycle của người thợ giỏi được lưu lại để đào tạo người mới.
Chúng em nói rõ: hiệu quả vận hành tại DENSO như thời gian time study dưới 30 phút là ước tính và cần pilot để xác nhận, theo đúng ba bước xác minh bên phải.
""")

    # 13 ── Kế hoạch triển khai (slide mẫu 7)
    s = clone(prs, T[7])
    set_text(find(s, "TextBox 6"), "Kế hoạch triển khai")
    for grp, t_name, b_name, lab, title, body in [
        ("Group 11", "TextBox 14", "TextBox 15", "ĐÃ XONG", "PoC hackathon",
         "Pipeline end-to-end trên HA4M, đánh giá vs nhãn, validator LLM, dashboard + video demo."),
        ("Group 23", "TextBox 26", "TextBox 27", "TUẦN 1–4", "Pilot 1 trạm",
         "Hiệu chỉnh camera, vẽ zone theo layout thật, đổi px→cm, đối chiếu với kỹ sư IE."),
        ("Group 35", "TextBox 38", "TextBox 39", "THÁNG 2–3", "Mở rộng 3–5 trạm",
         "Thư viện golden cycle, nạp SOP nội bộ vào KB, đo tỷ lệ đề xuất được kỹ sư chấp nhận."),
        ("Group 47", "TextBox 50", "TextBox 51", "THÁNG 4–6", "Nâng cấp model",
         "Phân đoạn học sâu (MS-TCN/ST-GCN) trên dữ liệu tích lũy, đa camera, đào tạo người mới."),
    ]:
        set_text(find(s, grp), lab)
        set_text(find(s, t_name), title)
        set_text(find(s, b_name), body)
    set_text(find(s, "Group 54"), "DELIVERABLES: dashboard theo trạm • baseline 3M • báo cáo có dẫn nguồn • golden "
                                  "cycle • backlog Kaizen")
    set_text(find(s, "TextBox 57"), "Giả định: video và SOP được cấp quyền truy cập; chỉ lưu keypoint để bảo vệ quyền "
                                    "riêng tư; mốc thời gian điều chỉnh theo phạm vi thực tế.")
    notes(s, """
[~1 phút — kế hoạch]
Giai đoạn hackathon đã xong: pipeline chạy end-to-end, có đánh giá độ chính xác, validator LLM và dashboard.
Bốn tuần đầu: pilot tại một trạm thật của DENSO – hiệu chỉnh camera, vẽ vùng khay theo layout thật, đổi đơn vị pixel sang centimet và đối chiếu với kỹ sư IE.
Tháng 2 đến 3: mở rộng 3 đến 5 trạm, xây thư viện golden cycle, nạp SOP nội bộ và đo tỷ lệ đề xuất được kỹ sư chấp nhận.
Tháng 4 đến 6: dùng dữ liệu tích lũy để huấn luyện model phân đoạn học sâu, hỗ trợ nhiều camera và đưa vào đào tạo công nhân mới.
Về quyền riêng tư, hệ thống chỉ cần lưu khung xương, không cần lưu khuôn mặt.
""")

    # 14 ── Kết luận (slide mẫu 8)
    s = clone(prs, T[8])
    set_text(find(s, "Group 7"), "KHÔNG CHỈ BẤM GIỜ — BIẾN THAO TÁC THÀNH TRI THỨC ĐO ĐƯỢC.")
    for a, b, c, t1, t2, t3 in [
        ("TextBox 12", "TextBox 13", "TextBox 14", "MEASURE", "Computer Vision",
         "Bước SOP, cycle time, 3M từ video camera thường."),
        ("TextBox 17", "TextBox 18", "TextBox 19", "EXPLAIN", "Kaizen Copilot",
         "LLM giải thích có trích dẫn, không được tạo số."),
        ("TextBox 22", "TextBox 23", "TextBox 24", "VERIFY", "Human-in-the-loop",
         "Kỹ sư duyệt; đo lại trước–sau cùng phương pháp."),
    ]:
        set_text(find(s, a), t1)
        set_text(find(s, b), t2)
        set_text(find(s, c), t3)
    set_text(find(s, "Group 27"), "Đề xuất: pilot 4 tuần tại một trạm DENSO — quay video, dashboard 3M có bằng chứng, "
                                  "kỹ sư duyệt và đo lại trước–sau.")
    set_text(find(s, "TextBox 30"), "Đo đúng • Giải thích có căn cứ • Kỹ sư quyết định • Tri thức được tích lũy")
    notes(s, """
[~45 giây — kết luận]
Tóm lại, AI Kaizen Copilot không chỉ bấm giờ: nó biến mỗi thao tác thành tri thức đo được.
MEASURE: đo bước SOP, cycle time và 3M từ camera thường.
EXPLAIN: Copilot giải thích có trích dẫn và không được phép tạo số liệu.
VERIFY: kỹ sư là người quyết định, và cải tiến được đo lại bằng đúng phương pháp đó.
Đề xuất của đội: pilot 4 tuần tại một trạm DENSO.
""")

    # 15 ── Kết thúc (slide mẫu 9)
    s = clone(prs, T[9])
    set_text(find(s, "TextBox 7"), "CẢM ƠN — HỎI & ĐÁP")
    set_text(find(s, "TextBox 8"), "DENSO Factory Hacks 2026  •  MEASURE · EXPLAIN · VERIFY")
    tb(s, 1.18, 10.5, 16.0, 0.4, "Trịnh Gia Huy (nhóm trưởng) · Vũ Văn Mạnh · Nguyễn Trần Quang · Trịnh Hoài Nam",
       size=14, color="FFFFFF")
    notes(s, """
[Hỏi & đáp]
Cảm ơn ban giám khảo đã lắng nghe. Đội sẵn sàng demo trực tiếp dashboard và trả lời câu hỏi.
Câu hỏi có thể gặp: (1) Khác gì bấm giờ bằng video thường? → đo tự động 100% chu kỳ, đủ 3M, có bằng chứng và đánh giá với nhãn chuyên gia.
(2) LLM có bịa không? → mọi số do code tính; validator chặn số/trích dẫn không có thật, có retry và báo tỷ lệ đạt.
(3) Áp dụng trạm khác thế nào? → chỉ cần vẽ zone và khai báo SOP cho trạm; mỗi lần chạy là một run để so sánh.
(4) Quyền riêng tư? → chỉ cần lưu khung xương; mục tiêu cải tiến quy trình, không giám sát cá nhân.
""")

    # Xoá 9 slide mẫu gốc (chỉ giữ slide mới)
    sld = prs.slides._sldIdLst
    for sid in list(sld)[:n_template]:
        prs.part.drop_rel(sid.rId)
        sld.remove(sid)
    prs.save(str(OUT))
    print(f"đã ghi {OUT} ({len(prs.slides)} slide)")
    write_script(prs)


SLIDE_TITLES = ["Bìa", "Tóm tắt đề xuất", "Giới thiệu team", "Bài toán tại line", "Luồng giải pháp",
                "Kiến trúc pipeline", "Dữ liệu & độ tin cậy keypoint", "Phân đoạn bước: CV vs nhãn",
                "Muda · Mura · Muri", "Findings & Kaizen Copilot", "Dashboard demo", "Hiệu quả mang lại",
                "Kế hoạch triển khai", "Kết luận", "Cảm ơn — Hỏi & đáp"]


def write_script(prs):
    """script.md = kịch bản thuyết trình, sinh từ CHÍNH speaker notes của deck (luôn khớp nhau)."""
    md = ["# Kịch bản thuyết trình — AI Kaizen Copilot · Motion Analyzer", "",
          f"Deck: `{OUT.name}` · {len(prs.slides)} slide · thời lượng mục tiêu ≈ 14–15 phút + hỏi đáp.",
          "Kịch bản dưới đây trùng với speaker notes trong file PowerPoint (sinh tự động bởi `build_slides.py`).",
          "", "Gợi ý phân vai: nhóm trưởng Trịnh Gia Huy mở đầu (slide 1–3) và kết thúc (slide 14–15); các phần "
          "giữa chia theo mảng mỗi thành viên phụ trách — đội tự điều chỉnh.", "",
          "Chuẩn bị demo: `streamlit run app/ui.py`, mở run "
          "`20261013-003627_pred_mediapipe_dry_u001-2trials` (link `?run=...`), tab *So sánh video*.", ""]
    for i, (slide, title) in enumerate(zip(prs.slides, SLIDE_TITLES), start=1):
        md += [f"## Slide {i} — {title}", "", slide.notes_slide.notes_text_frame.text.strip(), ""]
    (HERE / "script.md").write_text("\n".join(md), encoding="utf-8")
    print(f"đã ghi {HERE / 'script.md'}")


class _CellProxy:
    """Cho phép set_text() làm việc với ô bảng (có text_frame)."""

    def __init__(self, cell):
        self.text_frame = cell.text_frame


if __name__ == "__main__":
    build()
